import os
import json
import pytest
import re
import tempfile
import openpyxl
from datetime import date
from apps.academico.models import Aluno, Modalidade, Turma, Matricula, Categoria
from apps.localizacao.models import Polo
from apps.academico.views import _processar_registro_aluno, _persistir_alunos
from unittest.mock import patch
from django.urls import reverse
from django.contrib.auth.models import Permission, Group
from django.core.files.uploadedfile import SimpleUploadedFile
from playwright.sync_api import expect
    
@pytest.mark.django_db
def test_processar_registro_aluno_novo_com_faker(faker):
    
    # Geramos dados dinâmicos a cada execução do teste
    nome_falso = faker.name()
    data_falsa = faker.date_of_birth(minimum_age=18, maximum_age=60).strftime('%Y-%m-%d')
    
    daluno = {'nome': nome_falso, 'nascimento': data_falsa}
    metricas = {'estado_anterior': [], 'ids_criados': []}

    aluno = _processar_registro_aluno(daluno, metricas)

    # Verifica usando as variáveis geradas
    assert aluno.id is not None
    assert aluno.primeiro_nome == nome_falso.split()[0]
    
@pytest.mark.django_db
def test_persistir_alunos_fluxo_completo(faker, django_user_model):
    """
    Testa a persistência de alunos garantindo a criação de novos registros,
    atualização de existentes e criação de matrículas com dados dinâmicos.
    """
    # 1. SETUP DE AMBIENTE (Dados de Referência)
    user = django_user_model.objects.create_user(username=faker.user_name(), password='123')
    user.user_permissions.add(Permission.objects.get(codename='add_turma'))
    
    modalidade = Modalidade.objects.create(nome="Futsal")
    polo = Polo.objects.create(nome='FUNEL')
    cat = Categoria.objects.create(nome='Adulto')
    
    turma = Turma.objects.create(
        modalidade_id=modalidade,
        polo_id=polo,
        horario="14:00:00",
        categoria=cat,
        capacidade=30
    )

    # Criamos um aluno que já existe para testar a atualização (Update)
    aluno_existente = Aluno.objects.create(
        primeiro_nome="Maria",
        ultimo_nome="Silva",
        data_nascimento="1990-01-01",
        endereco=faker.address(),
        telefone=faker.msisdn()
    )

    # 2. PREPARAÇÃO DO PAYLOAD (Simulando dados do Excel via Faker)
    # Geramos um aluno novo e um para atualizar
    nome_novo = faker.name()
    nascimento_novo = faker.date_of_birth(minimum_age=10, maximum_age=50).strftime('%Y-%m-%d')
    
    alunos_dados = [
        {
            'id': None, # Novo aluno
            'nome': nome_novo,
            'nascimento': nascimento_novo,
            'status': 'NOVO',
            'data_atestado': faker.date_this_year().strftime('%Y-%m-%d')
        },
        {
            'id': aluno_existente.id, # Aluno para atualizar
            'nome': "Maria Silva Sauro", # Nome alterado
            'nascimento': "1990-01-01",
            'status': 'EXISTENTE',
            'data_atestado': "2026-05-14"
        }
    ]

    # 3. EXECUÇÃO DA FUNÇÃO
    metricas = _persistir_alunos(alunos_dados, turma, user)

    # 4. VALIDAÇÕES (ASSERTS)
    
    # Validação de Aluno Novo
    aluno_novo = Aluno.objects.get(primeiro_nome=nome_novo.split()[0])
    assert aluno_novo.id in metricas['ids_criados']
    assert metricas['qtd_matriculas'] >= 1
    
    # Validação de Aluno Atualizado
    aluno_existente.refresh_from_db()
    assert aluno_existente.ultimo_nome == "Silva Sauro"
    assert len(metricas['estado_anterior']) == 1
    assert metricas['estado_anterior'][0]['id'] == aluno_existente.id

    # Validação de Matrícula
    assert Matricula.objects.filter(aluno_id=aluno_novo, turma_id=turma).exists()
    assert Matricula.objects.filter(aluno_id=aluno_existente, turma_id=turma).exists()
    
@pytest.mark.django_db
def test_importar_ficha_fase2_salvamento(client, django_user_model):
    """ Testa a Fase 2: Recebendo o JSON do preview e gravando no banco """
    
    # 1. Setup do Usuário (Usando o is_staff para garantir acesso à área restrita)
    user = django_user_model.objects.create_user(username='admin', password='123', is_staff=True, is_superuser=True)
    permissao = Permission.objects.get(codename='add_turma')
    user.user_permissions.add(permissao)
    grupo_admin, _ = Group.objects.get_or_create(name='Secretaria')
    user.groups.add(grupo_admin)
    client.force_login(user)
    
    # 2. Inicialização das Tabelas de Referência
    Modalidade.objects.create(nome='FUTSAL')
    Polo.objects.create(nome='FUNEL')
    cat = Categoria.objects.create(nome='Adulto')

    # 3. Payload JSON (O Sinal Simulador)
    payload = {
        'turma': {
            'modalidade': 'FUTSAL',
            'local': 'FUNEL',
            'horario': '14:00',
            'categoria': 'Adulto',
            'dias': 'SEG, QUA',
            'professor': '',
            'estagiario': ''
        },
        'alunos': [
            {'nome': 'Lucas Teste', 'nascimento': '2000-01-01', 'atestado': '2026-05-20', 'status': 'NOVO'}
        ]
    }

    # 4. Disparo do POST
    url = reverse('academico:importar_ficha_chamada')
    response = client.post(url, {
        'confirmar_importacao': 'true',
        'payload_json': json.dumps(payload)
    })

    # 5. Verifica redirecionamento de sucesso
    assert response.status_code == 302
    assert response.url == reverse('academico:listar_turmas')

@pytest.mark.django_db
@patch('apps.academico.views.ler_ficha_chamada')
@patch('apps.academico.views.openpyxl.load_workbook')
def test_importar_ficha_fase1_preview_sucesso(mock_load_workbook, mock_ler_ficha, client, django_user_model, setup_dados):
    """ Testa se o envio de um arquivo Excel gera o HTML de preview corretamente """
    
    user = django_user_model.objects.create_user(username='admin', password='123', is_staff=True, is_superuser=True)
    user.user_permissions.add(Permission.objects.get(codename='add_turma'))
    grupo_admin, _ = Group.objects.get_or_create(name='Secretaria')
    user.groups.add(grupo_admin)
    client.force_login(user)
    
    professor = setup_dados['professor']

    Modalidade.objects.create(nome='FUTSAL')
    Polo.objects.create(nome='FUNEL')
    Categoria.objects.create(nome='ADULTO')

    # Mock da extração de dados
    mock_ler_ficha.return_value = {
        'modalidade': 'FUTSAL',
        'horario': '14:00',
        'categoria': 'ADULTO',
        'local': 'FUNEL',
        'dias': 'SEG/QUA',
        'professor': professor,
        'estagiario': '',
        'alunos': [{'nome': 'João Silva', 'nascimento': date(2000, 1, 1), 'atestado': None}]
    }

    # Arquivo binário simulado na memória RAM
    arquivo_simulado = SimpleUploadedFile("planilha_teste.xlsx", b"conteudo_falso", content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    url = reverse('academico:importar_ficha_chamada')
    response = client.post(url, {'arquivo_ficha': arquivo_simulado})

    # Verificações de Renderização
    assert response.status_code == 200
    assert 'academico/preview_importacao.html' in [t.name for t in response.templates]
    assert 'payload_json' in response.context
    
    payload = json.loads(response.context['payload_json'])
    assert payload['turma']['modalidade'] == 'FUTSAL'

@pytest.mark.django_db
def test_importar_ficha_fase2_erro_payload_vazio(client, django_user_model):
    """ Testa intertravamento: view bloqueia confirmação sem JSON """
    
    user = django_user_model.objects.create_user(username='admin', password='123', is_staff=True, is_superuser=True)
    user.user_permissions.add(Permission.objects.get(codename='add_turma'))
    
    grupo_admin, _ = Group.objects.get_or_create(name='Secretaria')
    user.groups.add(grupo_admin)
    client.force_login(user)

    url = reverse('academico:importar_ficha_chamada')
    response = client.post(url, {'payload_json': ''})

    assert response.status_code == 302
    assert response.url == url

@pytest.mark.django_db
@patch('apps.academico.views.ler_ficha_chamada')
@patch('apps.academico.views.openpyxl.load_workbook')
def test_importar_ficha_fase1_erro_modalidade_nula(mock_load_workbook, mock_ler_ficha, client, django_user_model):
    """ Testa bloqueio de planilha com dados vitais corrompidos """
    
    user = django_user_model.objects.create_user(username='admin', password='123', is_staff=True, is_superuser=True)
    user.user_permissions.add(Permission.objects.get(codename='add_turma'))
    
    grupo_admin, _ = Group.objects.get_or_create(name='Secretaria')
    user.groups.add(grupo_admin)
    client.force_login(user)

    mock_ler_ficha.return_value = {
        'modalidade': None, 
        'horario': '14:00',
        'local': 'FUNEL',
        'categoria': 'Adulto',
        'dias': 'SEG/QUA',
        'alunos': []
    }

    arquivo_simulado = SimpleUploadedFile("planilha.xlsx", b"dados", content_type="application/octet-stream")
    url = reverse('academico:importar_ficha_chamada')
    response = client.post(url, {'arquivo_ficha': arquivo_simulado})

    assert response.status_code == 302
    assert response.url == url
    
@pytest.mark.django_db
def test_importacao_falha_sem_data_nascimento_ou_atestado(client, django_user_model):
    """
    Testa o Intertravamento do Banco: 
    Envia um payload com dados faltantes e verifica se o sistema aborta a transação.
    """
    user = django_user_model.objects.create_user(username='admin', password='123', is_staff=True, is_superuser=True)
    user.user_permissions.add(Permission.objects.get(codename='add_turma'))
    
    grupo_admin, _ = Group.objects.get_or_create(name='Secretaria')
    user.groups.add(grupo_admin)
    client.force_login(user)
    
    url = reverse('academico:importar_ficha_chamada')
    
    # Pacote com a anomalia (nascimento = '-')
    payload_corrompido = {
        "turma": {
            "modalidade": "VÔLEI", "local": "FUNEL", "horario": "14:00", 
            "dias": "SEGUNDA, QUARTA", "capacidade": 20, 'categoria': 'Adulto'
        },
        "alunos": [
            {
                "nome": "ALUNO DEFEITUOSO",
                "nascimento": "-",  # <--- Curto-circuito intencional
                "atestado": "2026-05-20",
                "status": "NOVO"
            }
        ]
    }
    
    
    response = client.post(url, {'payload_json': json.dumps(payload_corrompido)})
    
    # O sistema não pode dar erro 500, ele deve capturar o ValueError 
    # e fazer o redirecionamento (302) de volta para a tela de upload.
    assert response.status_code == 302
    assert response.url == reverse('academico:importar_ficha_chamada')

# ==========================================
# BLOCO 2: TESTES END-TO-END (PLAYWRIGHT)
# ==========================================

# TRAVA CRÍTICA: O transaction=True é obrigatório para o live_server enxergar o banco
@pytest.mark.django_db(transaction=True)
def test_fluxo_importacao_na_tela_com_playwright(page, live_server, django_user_model, client, setup_dados):
    """ E2E: Simula operação robótica completa da IHM """
    
    # 1. Setup da Matriz Base
    user = django_user_model.objects.create_user(username='admin', password='123', is_staff=True, is_superuser=True)
    user.user_permissions.add(Permission.objects.get(codename='add_turma'))
    grupo_admin, _ = Group.objects.get_or_create(name='Secretaria')
    user.groups.add(grupo_admin)
    
    professor = setup_dados['professor']

    # 3. Força o Login
    client.force_login(user)
    Modalidade.objects.create(nome='FUTSAL')
    Polo.objects.create(nome='FUNEL')
    Categoria.objects.create(nome='ADULTO')

    # 2. Gerador de Arquivo Estéril (Para rodar no GitHub Actions sem falhar)
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # Cabeçalhos do Eixo Y (A âncora na coluna A, o valor na coluna B)
    ws['A1'] = 'Modalidade'
    ws['B1'] = 'FUTSAL'
    ws['C1'] = 'Local'
    ws['D1'] = 'FUNEL'
    ws['E2'] = 'Categoria'
    ws['F2'] = 'ADULTO'
    
    ws['A2'] = 'Horário'
    ws['B2'] = '14:00'
    ws['C2'] = 'Professor'
    ws['D2'] = str(professor)
    
    # O campo de dias usa a marcha 'mesma_celula'
    ws['E1'] = 'Dias: SEG/QUA' 

    # Eixo X (Tabela de Alunos) - Cria as âncoras da tabela
    ws['A4'] = 'Nome do Aluno'
    ws['B4'] = 'Data Nasc.'
    ws['C4'] = 'Validade Atestado'
    
    # Inserção de dados dos alunos logo abaixo
    ws['A5'] = 'Aluno Automatizado'
    ws['B5'] = '2001-04-12'
    ws['C5'] = '2026-04-03'
    
    # Salva na pasta temporária do sistema operacional
    fd, caminho_planilha = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd) # Libera o pino do arquivo para escrita
    wb.save(caminho_planilha)

    try:
        # 3. Execução Autônoma (Login)
        page.goto(f"{live_server.url}/accounts/login/")
        page.fill("input[name='username']", "admin")
        page.fill("input[name='password']", "123")
        with page.expect_navigation():
            page.click("button[type='submit']")

        # 4. Fase de Upload
        url_importacao = f"{live_server.url}{reverse('academico:importar_ficha_chamada')}"
        page.goto(url_importacao)
        
        # Injeta o arquivo recém-fabricado no input file
        page.set_input_files("input[type='file'][name='arquivo_ficha']", caminho_planilha)
        
        with page.expect_navigation():
            page.click("button:has-text('Carregar e Visualizar Dados')")

        # 5. Validação Visual
        page.wait_for_selector("text=FUTSAL") 
        page.click("button[name='confirmar_importacao']")

        # 6. Telemetria Final
        page.wait_for_url(f"{live_server.url}{reverse('academico:listar_turmas')}")
        assert "Sucesso" in page.inner_text("body")
        
    finally:
        # 7. Limpeza da Bancada (Rotina de Destruição do Arquivo Temp)
        if os.path.exists(caminho_planilha):
            os.remove(caminho_planilha)

def test_ihm_rele_sobrecapacidade(page, live_server, django_user_model, client, setup_dados):
    """
    Testa o Sensor de Capacidade:
    Importa 3 alunos, tenta setar a capacidade para 2 e clica em salvar.
    """
    
    user = django_user_model.objects.create_user(username='admin', password='123', is_staff=True, is_superuser=True)
    user.user_permissions.add(Permission.objects.get(codename='add_turma'))
    
    grupo_admin, _ = Group.objects.get_or_create(name='Secretaria')
    user.groups.add(grupo_admin)
    
    professor = setup_dados['professor']
    
    client.force_login(user)
    
    Modalidade.objects.create(nome='FUTSAL')
    Polo.objects.create(nome='FUNEL')
    Categoria.objects.create(nome='ADULTO')
    
    # Gerando planilha com dados sintéticos
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'Modalidade'; ws['B1'] = 'FUTSAL'
    ws['C1'] = 'Local';      ws['D1'] = 'FUNEL'
    ws['A2'] = 'Horário';    ws['B2'] = '14:00'
    ws['C2'] = 'Professor';  ws['D2'] = str(professor)
    ws['E1'] = 'Dias: SEG/QUA'; 
    ws['E2'] = 'Categoria'; ws['F2'] = 'ADULTO'
    
    ws['A4'] = 'Nº'
    ws['B4'] = 'Nome do Aluno'
    ws['C4'] = 'Data Nascimento'
    ws['D4'] = 'Atestado'
    
    linhas = [5, 6, 7]
    for i, linha in enumerate(linhas):
        ws[f'A{linha}'] = f'{i+1}'
        ws[f'B{linha}'] = f'Aluno Teste {i+1}'
        ws[f'C{linha}'] = '2001-04-12'
        ws[f'D{linha}'] = '2026-04-03'
    
    fd, caminho_planilha = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd) # Libera o pino do arquivo para escrita
    wb.save(caminho_planilha)
    
    try:
        page.goto(f"{live_server.url}/accounts/login/")
        page.fill("input[name='username']", "admin")
        page.fill("input[name='password']", "123")
        with page.expect_navigation():
            page.click("button[type='submit']")
        
        page.goto(f"{live_server.url}{reverse('academico:importar_ficha_chamada')}")
        
        # 1. Faz o upload da planilha
        page.set_input_files("input[type='file'][name='arquivo_ficha']", caminho_planilha)
        with page.expect_navigation():
            page.click("button:has-text('Carregar e Visualizar Dados')")

        # 2. Configura o robô para capturar o alerta (A Sirene do Sistema)
        alerta_disparado = []
        
        def handle_dialog(dialog):
            alerta_disparado.append(dialog.message)
            dialog.accept() # O robô aperta "OK" no alerta para não travar a esteira!
            
        page.on("dialog", handle_dialog)

        # 3. Manipula a capacidade para forçar o erro
        page.fill("#edit-capacidade", "2")
        
        # 4. Tenta salvar (o disparo deve ocorrer aqui)
        page.click("button[name='confirmar_importacao']")

        # 5. Telemetria de Validação
        assert len(alerta_disparado) > 0, "O disjuntor de capacidade falhou em disparar o alerta."
        assert "Limite Físico" in alerta_disparado[0]
        
        # Verifica se o LED vermelho acendeu no input
        expect(page.locator("#edit-capacidade")).to_have_class(re.compile(r"text-danger"))
    finally:
        # 7. Limpeza da Bancada (Rotina de Destruição do Arquivo Temp)
        if os.path.exists(caminho_planilha):
            os.remove(caminho_planilha)


def test_ihm_led_vermelho_data_ausente(page, live_server, django_user_model, client):
    """
    Testa o Alarme Visual:
    Garante que a IHM renderiza a borda vermelha quando um aluno vem sem data do Excel.
    """
    
    user = django_user_model.objects.create_user(username='admin', password='123', is_staff=True, is_superuser=True)
    user.user_permissions.add(Permission.objects.get(codename='add_turma'))
    grupo_admin, _ = Group.objects.get_or_create(name='Secretaria')
    user.groups.add(grupo_admin)
    
    prof = django_user_model.objects.create_user(username='prof', password='123', first_name='Bruno', last_name='Alves', is_staff=True, is_superuser=True)
    grupo_prof, _ = Group.objects.get_or_create(name='Professor')
    prof.groups.add(grupo_prof)
    
    Polo.objects.get_or_create(id='1', nome='FUNEL')
    Modalidade.objects.get_or_create(nome="Futsal")
    Categoria.objects.create(nome='ADULTO')
    
    client.force_login(user)
    
    # Gerando planilha com dados sintéticos
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'Modalidade'; ws['B1'] = 'FUTSAL'
    ws['C1'] = 'Local';      ws['D1'] = 'FUNEL'
    ws['A2'] = 'Horário';    ws['B2'] = '14:00'
    ws['C2'] = 'Professor';  ws['D2'] = 'Bruno Alves'
    ws['E1'] = 'Dias: SEG/QUA'
    ws['E2'] = 'Categoria'; ws['F2'] = 'ADULTO'

    ws['A4'] = 'Nº'
    ws['B4'] = 'Nome do Aluno'
    ws['C4'] = 'Data Nascimento'
    ws['D4'] = 'Validade Atestado'

    # Aluno 1: Carga Perfeita
    ws['A5'] = '1'
    ws['B5'] = 'Aluno Perfeito'
    ws['C5'] = '2001-04-12'
    ws['D5'] = '2026-04-03'

    # Aluno 2: Curto-circuito intencional (Falta Data de Nascimento)
    ws['A6'] = '2'
    ws['B6'] = 'Aluno Defeituoso'
    ws['C6'] = '-' # Coluna de nascimento vazia ou traco
    ws['D6'] = '2026-04-03'
    
    fd, caminho_planilha = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd) # Libera o pino do arquivo para escrita
    wb.save(caminho_planilha)
    
    try:
        page.goto(f"{live_server.url}/accounts/login/")
        page.fill("input[name='username']", "admin")
        page.fill("input[name='password']", "123")
        with page.expect_navigation():
            page.click("button[type='submit']")
        
        page.goto(f"{live_server.url}{reverse('academico:importar_ficha_chamada')}")
        
        # 1. Faz o upload da planilha
        page.set_input_files("input[type='file'][name='arquivo_ficha']", caminho_planilha)
        with page.expect_navigation():
            page.click("button:has-text('Carregar e Visualizar Dados')")
        
        # 1. Aguarda a tabela renderizar
        page.wait_for_selector("#tbody-alunos")
        
        # 2. Verifica se existe pelo menos um input com a classe de erro 'border-danger'
        input_com_defeito = page.locator("input.border-danger").first
        
        # O robô deve confirmar que o LED está aceso na tela
        expect(input_com_defeito).to_be_visible()
    finally:
        # 7. Limpeza da Bancada (Rotina de Destruição do Arquivo Temp)
        if os.path.exists(caminho_planilha):
            os.remove(caminho_planilha)